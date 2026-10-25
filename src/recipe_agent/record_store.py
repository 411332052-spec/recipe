"""Excel 紀錄：永遠往下加一列，不覆蓋。寫不進去（例如你正開著 Excel）就印警告、重試，
還是不行就先存進 pending_records.jsonl，下次寫入時補寫。
"""

from __future__ import annotations

import json
import logging
import os
import time
import zipfile
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

log = logging.getLogger("recipe_agent.record")

HEADERS = [
    "日期時間",
    "上傳者",
    "是否為食材",
    "辨識到的食材清單",
    "推薦料理名稱",
    "完整食譜與步驟",
    "照片存放位置",
    "備註",
]
_WIDTHS = [20, 30, 12, 30, 20, 70, 45, 40]
_MAX_CELL = 32000  # Excel 單格上限約 32767 字
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


@dataclass
class Record:
    when: str
    uploader: str
    is_food: str
    ingredients: str
    dish_name: str
    recipe: str
    photo_files: str
    note: str = ""

    def as_row(self) -> list[str]:
        return [getattr(self, f.name) for f in fields(self)]


class RecordFileProblem(Exception):
    """recipes.xlsx 壞了或打不開（不是被鎖住）。絕不覆蓋它。"""


class RecordStore:
    def __init__(
        self,
        excel_path: Path,
        pending_path: Path,
        retries: int = 3,
        retry_wait: float = 5.0,
        sleep: Callable[[float], None] = time.sleep,
    ) -> None:
        self.excel_path = excel_path
        self.pending_path = pending_path
        self.retries = retries
        self.retry_wait = retry_wait
        self._sleep = sleep

    # ---------- 對外 ----------

    def append(self, record: Record) -> bool:
        """加一列。回傳 True＝已寫進 Excel；False＝先暫存到 pending（之後補寫）。永遠不丟例外。"""
        try:
            return self._append(record)
        except Exception:
            log.exception("寫入紀錄時發生預料外的錯誤，改存暫存清單")
            return self._stash([record])

    def check_readable(self) -> None:
        """啟動時呼叫：recipes.xlsx 已存在卻打不開（壞了）就丟 RecordFileProblem，絕不覆蓋它。"""
        if not self.excel_path.exists():
            return
        try:
            load_workbook(self.excel_path).close()
        except PermissionError:
            return  # 只是正被 Excel 開著，不是壞掉；之後寫入時會重試
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise RecordFileProblem(str(exc)) from exc

    def flush_pending(self) -> None:
        """把暫存清單補寫進 Excel（啟動時呼叫）。失敗就保留，不丟例外。"""
        try:
            rows = self._read_pending()
            if rows and self._try_write(rows):
                self._clear_pending()
                log.info("已把 %d 筆暫存紀錄補寫進 Excel", len(rows))
        except Exception:
            log.exception("補寫暫存紀錄失敗，暫存清單保留")

    # ---------- 內部 ----------

    def _append(self, record: Record) -> bool:
        pending = self._read_pending()
        rows = [*pending, record]
        for attempt in range(1, self.retries + 1):
            ok = self._try_write(rows, warn=True, attempt=attempt)
            if ok:
                if pending:
                    self._clear_pending()
                    log.info("已把 %d 筆暫存紀錄補寫進 Excel", len(pending))
                return True
            if attempt < self.retries:
                self._sleep(self.retry_wait)
        print(f"[警告] Excel 還是寫不進去，這一筆先存到 {self.pending_path.name}，下次寫入時補寫。", flush=True)
        log.warning("Excel 寫入失敗，改存暫存清單")
        # pending 裡原本就有的不重複存，只補這一筆
        return self._stash([record])

    def _try_write(self, rows: list[Record], warn: bool = False, attempt: int = 0) -> bool:
        try:
            self._write_rows(rows)
            return True
        except PermissionError:
            msg = f"Excel 檔被鎖住了（你是不是正開著 {self.excel_path.name}？）"
            if warn:
                print(f"[警告] {msg}，第 {attempt} 次寫入失敗。", flush=True)
            log.warning("%s attempt=%s", msg, attempt)
            return False
        except RecordFileProblem:
            log.error("recipes.xlsx 無法讀取，為避免覆蓋，不會動它")
            return False

    def _write_rows(self, rows: list[Record]) -> None:
        if self.excel_path.exists():
            try:
                wb = load_workbook(self.excel_path)
            except (zipfile.BadZipFile, KeyError, ValueError) as exc:
                raise RecordFileProblem(str(exc)) from exc
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "紀錄"
            ws.append(HEADERS)
            for col, width in enumerate(_WIDTHS, start=1):
                ws.column_dimensions[chr(64 + col)].width = width
                ws.cell(row=1, column=col).font = Font(bold=True)
            ws.freeze_panes = "A2"

        for rec in rows:
            r = ws.max_row + 1
            for col, value in enumerate(rec.as_row(), start=1):
                text = ILLEGAL_CHARACTERS_RE.sub("", str(value))[:_MAX_CELL]
                cell = ws.cell(row=r, column=col, value=text)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

                # ==== 防護 4：Excel 文字轉純文字　開始 ====
                if text.startswith(_FORMULA_PREFIXES):
                    cell.value = "'" + text  # 開頭加單引號，Excel 就不會把它當公式
                    cell.data_type = "s"  # 同時明確標成文字
                # ==== 防護 4 結束 ====

        tmp = self.excel_path.with_suffix(".xlsx.tmp")
        wb.save(tmp)  # 先存暫存檔，成功才取代正式檔，當機也不會弄壞整份紀錄
        os.replace(tmp, self.excel_path)

    def _stash(self, records: list[Record]) -> bool:
        self.pending_path.parent.mkdir(parents=True, exist_ok=True)
        with open(self.pending_path, "a", encoding="utf-8") as fh:
            for rec in records:
                fh.write(json.dumps(asdict(rec), ensure_ascii=False) + "\n")
        return False

    def _read_pending(self) -> list[Record]:
        if not self.pending_path.exists():
            return []
        out: list[Record] = []
        for line in self.pending_path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                out.append(Record(**json.loads(line)))
            except (json.JSONDecodeError, TypeError):
                log.warning("暫存清單有一行讀不懂，已略過：%.80s", line)
        return out

    def _clear_pending(self) -> None:
        self.pending_path.unlink(missing_ok=True)
