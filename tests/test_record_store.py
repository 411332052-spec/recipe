import json

import pytest
from openpyxl import load_workbook

from recipe_agent.record_store import HEADERS, Record, RecordFileProblem, RecordStore


def _rec(**kw):
    base = dict(when="2026-10-25 18:30:05", uploader="小明 (1)", is_food="是", ingredients="番茄、雞蛋",
                dish_name="番茄炒蛋", recipe="1. 炒", photo_files="a.jpg")
    base.update(kw)
    return Record(**base)


def _store(tmp_path, **kw):
    kw.setdefault("sleep", lambda s: None)
    return RecordStore(tmp_path / "recipes.xlsx", tmp_path / "pending_records.jsonl", **kw)


def _rows(path):
    return [[c.value for c in row] for row in load_workbook(path).active.iter_rows()]


def test_creates_file_with_headers_and_appends(tmp_path):
    s = _store(tmp_path)
    assert s.append(_rec()) is True
    rows = _rows(s.excel_path)
    assert rows[0] == HEADERS and rows[1][4] == "番茄炒蛋"


def test_appends_never_overwrites(tmp_path):
    s = _store(tmp_path)
    s.append(_rec(dish_name="第一"))
    s2 = _store(tmp_path)  # 模擬重開程式
    s2.append(_rec(dish_name="第二"))
    s2.append(_rec(dish_name="第三"))
    assert [r[4] for r in _rows(s.excel_path)[1:]] == ["第一", "第二", "第三"]


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r", "\n"])
def test_guard4_formula_like_text_is_stored_as_text(tmp_path, prefix):
    s = _store(tmp_path)
    s.append(_rec(uploader=f"{prefix}1+1", recipe=f"{prefix}SUM(A1)"))
    wb = load_workbook(s.excel_path)
    for cell in (wb.active["B2"], wb.active["F2"]):
        assert cell.data_type == "s" and cell.value.startswith("'")


def test_guard4_normal_text_is_untouched(tmp_path):
    s = _store(tmp_path)
    s.append(_rec(dish_name="番茄炒蛋"))
    assert load_workbook(s.excel_path).active["E2"].value == "番茄炒蛋"


def test_illegal_control_chars_and_huge_text_do_not_crash(tmp_path):
    s = _store(tmp_path)
    assert s.append(_rec(recipe="a\x00b\x07c" + "字" * 40000)) is True
    v = load_workbook(s.excel_path).active["F2"].value
    assert v.startswith("abc") and len(v) <= 32001


def test_locked_excel_retries_warns_then_stashes_and_later_flushes(tmp_path, monkeypatch, capsys):
    s = _store(tmp_path, retries=3)
    s.append(_rec(dish_name="先寫一筆"))

    real = RecordStore._write_rows
    locked = {"on": True}

    def fake(self, rows):
        if locked["on"]:
            raise PermissionError("locked")
        return real(self, rows)

    monkeypatch.setattr(RecordStore, "_write_rows", fake)
    assert s.append(_rec(dish_name="被鎖住的")) is False
    out = capsys.readouterr().out
    assert out.count("[警告]") >= 3  # 每次重試都印警告
    pending = [json.loads(l) for l in s.pending_path.read_text(encoding="utf-8").splitlines()]
    assert [p["dish_name"] for p in pending] == ["被鎖住的"]

    locked["on"] = False
    assert s.append(_rec(dish_name="解鎖後")) is True
    assert [r[4] for r in _rows(s.excel_path)[1:]] == ["先寫一筆", "被鎖住的", "解鎖後"]
    assert not s.pending_path.exists()


def test_flush_pending_on_startup(tmp_path, monkeypatch):
    s = _store(tmp_path, retries=1)
    monkeypatch.setattr(RecordStore, "_write_rows", lambda self, rows: (_ for _ in ()).throw(PermissionError()))
    s.append(_rec(dish_name="暫存"))
    monkeypatch.undo()
    _store(tmp_path).flush_pending()
    assert [r[4] for r in _rows(s.excel_path)[1:]] == ["暫存"]


def test_corrupt_excel_is_not_overwritten_and_record_is_kept(tmp_path):
    s = _store(tmp_path, retries=1)
    s.excel_path.write_bytes(b"this is not a zip")
    assert s.append(_rec()) is False
    assert s.excel_path.read_bytes() == b"this is not a zip"
    assert s.pending_path.exists()


def test_append_never_raises_on_unexpected_error(tmp_path, monkeypatch):
    s = _store(tmp_path)
    monkeypatch.setattr(RecordStore, "_write_rows", lambda self, rows: (_ for _ in ()).throw(RuntimeError("boom")))
    assert s.append(_rec()) is False


def test_check_readable_ok_when_missing_or_good_and_raises_when_corrupt(tmp_path):
    s = _store(tmp_path)
    s.check_readable()  # 檔案還不存在：沒事
    s.append(_rec())
    s.check_readable()  # 正常檔案：沒事
    s.excel_path.write_bytes(b"not a zip")
    with pytest.raises(RecordFileProblem):
        s.check_readable()
    assert s.excel_path.read_bytes() == b"not a zip"
